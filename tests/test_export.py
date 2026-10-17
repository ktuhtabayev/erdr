"""Exporters: tables, text formats, Excel, figures, the HTML report and the datasets."""

import csv
import json
from pathlib import Path

import openpyxl
import pytest

from erdr.data import load_dataset
from erdr.errors import ConfigError
from erdr.export import EXPORTERS, ExportOptions, export_result, resolve_formats, run_settings
from erdr.export.figures import (
    RUN_FIGURES,
    classical_mds,
    figure_names,
    run_figure,
    save_figure,
    svg_text,
)
from erdr.export.tables import MATRIX_LIMIT, run_tables
from erdr.export.text import display, exact, latex_escape, to_csv, to_latex, to_markdown
from erdr.services.analysis import compare_readings, sweep_k
from erdr.services.experiment import ExperimentResult, run_experiment


@pytest.fixture(scope="module")
def options(default_run: ExperimentResult) -> ExportOptions:
    return ExportOptions(
        readings=compare_readings(default_run.dataset),
        sweep=sweep_k(default_run.dataset, ks=[1, 3]),
    )


def test_tables_mirror_the_workbook(default_run: ExperimentResult):
    tables = {t.key: t for t in run_tables(default_run)}
    assert list(tables) == [
        "summary",
        "dataset",
        "normalized",
        "distances",
        "distances-quantitative",
        "distances-nominal",
        "neighbours",
        "characteristics",
        "spheres",
        "enemy-spheres",
        "k-spheres",
        "covering",
        "sphere-checks",
        "queue",
        "iterations",
        "iteration-spheres",
        "sphere-trace",
        "counters",
        "routes",
        "composition",
        "cases",
    ]
    iterations = tables["iterations"]
    assert [row[1] for row in iterations.rows] == ["S₁", "S₃", "S₅", "S₂", "S₆", "S₄"]
    assert iterations.rows[2][12] == "S₁ (Step 5), S₅ (Step 5)"
    assert iterations.mark(0, 1) == "selected"
    routes = tables["routes"]
    assert [row[2] for row in routes.rows].count("D · training") == 5
    queue = tables["queue"]
    assert queue.rows[0][6] == "index (equal d and r)"
    assert queue.mark(1, 6) == "tie-r"
    assert tables["distances"].matrix
    assert tables["cases"].title.startswith("Case checklist · 17 of 17")
    assert tables["summary"].note == "Literal article reading."


def test_text_formats(default_run: ExperimentResult):
    table = next(t for t in run_tables(default_run) if t.key == "characteristics")
    rows = list(csv.reader(to_csv(table).splitlines()))
    assert rows[0][3] == "r(Sᵢ)"
    assert float(rows[1][3]) == default_run.result.characteristics.r[0]  # exact
    assert "| S₁ |" in to_markdown(table)
    latex = to_latex(table)
    assert latex.isascii()
    assert r"\begin{tabular}" in latex
    assert display(0.07954283880000001) == "0.0795428388"
    assert display(2.0) == "2"
    assert display(None) == ""
    assert exact(0.1) == "0.1"
    assert latex_escape("S₁ → ρ_I & ∅") == r"S$_{1}$ $\to$ $\rho$\_I \& $\emptyset$"
    assert latex_escape("Sᵢ→Sⱼ") == r"S$_{i}\to$S$_{j}$"  # adjacent math shares one $…$


def test_resolve_formats():
    assert resolve_formats(["all"]) == EXPORTERS.names()
    assert resolve_formats(["csv, excel", "csv"]) == ["csv", "xlsx"]
    with pytest.raises(ConfigError, match="unknown export format"):
        resolve_formats(["pdf"])


def test_export_everything(tmp_path: Path, default_run: ExperimentResult, options: ExportOptions):
    files = export_result(default_run, tmp_path, options=options)
    names = {p.relative_to(tmp_path).as_posix() for p in files}
    assert {"run.json", "report.md", "report.html", "workbook.xlsx"} <= names
    assert "csv/readings.csv" in names
    assert "figures/sweep.svg" in names
    assert "latex/tables.tex" in names
    assert {
        f"datasets/{part}.{fmt}" for part in ("training", "validation") for fmt in ("dat", "csv")
    } <= names
    training = load_dataset(tmp_path / "datasets" / "training.dat")
    assert training.m == 5
    data = json.loads((tmp_path / "run.json").read_text(encoding="utf-8"))
    assert data["meta"]["settings"]["Reading"] == "A=1 B=1 C=1 · z: any class"
    book = openpyxl.load_workbook(tmp_path / "workbook.xlsx")
    assert book.sheetnames[:3] == ["Overview", "Dataset", "Normalized Dataset"]
    assert "Charts" in book.sheetnames
    assert "Zhuravlev Distances" in book.sheetnames
    html = (tmp_path / "report.html").read_text(encoding="utf-8")
    assert html.count("<svg") == len(figure_names()) + 2
    assert "Not the literal article reading" not in html
    csv_bytes = (tmp_path / "csv" / "characteristics.csv").read_bytes()
    assert csv_bytes.startswith(b"\xef\xbb\xbf")


def test_large_samples_leave_matrices_out_of_text_reports(tmp_path: Path, heart270):
    run = run_experiment(dataset=heart270)
    assert run.result.m > MATRIX_LIMIT
    export_result(run, tmp_path, ["markdown", "html", "latex"])
    assert "Distances ρ" not in (tmp_path / "report.md").read_text(encoding="utf-8")
    assert "Left out here" in (tmp_path / "report.html").read_text(encoding="utf-8")
    assert not (tmp_path / "latex" / "distances.tex").exists()


def test_figures(tmp_path: Path, default_run: ExperimentResult):
    for key in RUN_FIGURES:
        figure = run_figure(key, default_run)
        assert figure.axes
    path = save_figure(run_figure("map", default_run), tmp_path / "map.svg")
    assert "<svg" in path.read_text(encoding="utf-8")
    assert svg_text(run_figure("sizes", default_run)).startswith("<svg")
    coords = classical_mds(default_run.distances.matrix)
    assert coords.shape == (7, 2)
    assert classical_mds(default_run.distances.matrix[:1, :1]).shape == (1, 2)


def test_settings_name_the_reading(heart270):
    from erdr.config import preset

    run = run_experiment(preset("closest"), heart270)
    settings = dict(run_settings(run))
    assert settings["Reading"] == "A=1 B=2 C=1 · z: own class"


def test_settings_name_the_dataset_type_and_flag_a_partial_metric(heart270):
    from erdr.config import ExperimentConfig, plugin
    from erdr.config.models import DistanceConfig

    settings = dict(run_settings(run_experiment(dataset=heart270)))
    assert settings["Dataset type"] == "heterogeneous"
    assert settings["Metric · features"] == "zhuravlev (auto · heterogeneous dataset) · auto"
    assert "⚠ Metric" not in settings
    euclidean = ExperimentConfig(distance=DistanceConfig(metric=plugin("euclidean")))
    settings = dict(run_settings(run_experiment(euclidean, heart270)))
    assert settings["Metric · features"] == "euclidean · auto"
    assert "7 nominal feature(s) do not count" in settings["⚠ Metric"]
