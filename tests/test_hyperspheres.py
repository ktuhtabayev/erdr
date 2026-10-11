"""Steps 1–5 as hyperspheres: the core, the tables, the figures, the exports and the validator."""

import math
import shutil
import warnings
from pathlib import Path

import numpy as np
import openpyxl
import pytest
from hypothesis import given, settings
from hypothesis import strategies as st
from PIL import Image

from erdr.core.algorithm import run_erdr
from erdr.core.cases import SHORT_TITLES, cases_in_iteration
from erdr.core.characteristics import compute_characteristics
from erdr.core.hyperspheres import (
    RATIO_LINES,
    Relation,
    Role,
    angle_step,
    coinciding_points,
    cover_change_text,
    density_text,
    iteration_spheres,
    iteration_title,
    object_title,
    possible_cover,
    priority_from_ratio,
    radial_layout,
    ratio_from_priority,
    ratio_line,
    ray_length,
    relation_meaning,
    sphere_profile,
    step1_summary,
    strip_levels,
)
from erdr.core.readings import CLOSEST, LITERAL, Readings, ZScope, all_readings
from erdr.export import ExportOptions, export_result
from erdr.export.figures import (
    CROWDED,
    FIGURE_SHEETS,
    PANEL_FIGURES,
    PANEL_LIMIT,
    figure_names,
    figure_sheet,
    iteration_animation,
    iteration_spheres_figure,
    object_spheres_figure,
    run_figure,
    sphere_iterations,
    sphere_objects,
)
from erdr.export.tables import (
    HYPERSPHERES_ITERATIONS,
    HYPERSPHERES_STEP1,
    covering_table,
    enemy_spheres_table,
    iteration_detail_table,
    iteration_spheres_table,
    k_spheres_table,
    object_spheres_table,
    run_tables,
    sphere_checks_table,
    sphere_trace_table,
)
from erdr.export.theme import stage_of
from erdr.services.experiment import ExperimentResult, run_experiment
from erdr.services.validation import validate_workbook
from tests.conftest import WORKBOOK
from tests.test_core_algorithm import samples

# The workbook's sheet *Hyperspheres (Step 1)* (S₁ … S₇)
OWN_INSIDE = [2, 3, 2, 0, 2, 0, 1]
CLOSED = [4, 5, 4, 2, 4, 2, 3]
K_OWN = [2, 3, 2, 0, 2, 2, 2]
K_ENEMIES = [1, 0, 1, 3, 1, 1, 1]
RELATIONS = ["EQUAL", "INSIDE", "EQUAL", "OUTSIDE", "EQUAL", "OUTSIDE", "OUTSIDE"]


# ---------------------------------------------------------------- the core: Step 1


def test_sphere_profile_reproduces_the_workbook(default_run: ExperimentResult):
    c = default_run.result.characteristics
    p = sphere_profile(default_run.distances.matrix, c)
    assert p.m == 7
    assert p.own_inside.tolist() == OWN_INSIDE
    assert p.own_on_sphere.tolist() == [0] * 7
    assert p.enemies_inside.tolist() == [0] * 7  # by the definition of r
    assert p.enemies_on_sphere.tolist() == [1] * 7  # the nearest enemy
    assert p.closed_count.tolist() == CLOSED
    assert p.k_members.tolist() == [3] * 7
    assert p.k_own.tolist() == K_OWN
    assert p.k_enemies.tolist() == K_ENEMIES
    assert [r.name for r in p.relation] == RELATIONS
    assert p.d_from_ratio.tolist() == pytest.approx(c.d.tolist(), abs=1e-12)
    assert np.array_equal(p.containing, c.spheres.T)
    assert p.spheres_containing(1) == (0, 1, 4)  # S₁, S₂ (itself), S₅
    assert p.objects_with(Relation.EQUAL) == (0, 2, 4)
    lines = step1_summary(c, p, default_run.object_names)
    assert [line.text for line in lines[:5]] == [
        "S₄, S₆",
        "S₁, S₂, S₃, S₅, S₇",
        "S₁, S₃, S₅",
        "S₂",
        "S₄, S₆, S₇",
    ]
    assert all(line.passed for line in lines[5:])


def test_ratio_and_priority_are_inverse():
    for d in (0.0, 0.1, 0.25, 0.9):
        assert priority_from_ratio(ratio_from_priority(d)) == pytest.approx(d)
    below, above = ratio_line(0.1, 10.0)
    assert below == pytest.approx((10.0, 9 / 11 * 10))
    assert above == pytest.approx((9 / 11 * 10, 10.0))
    assert RATIO_LINES[0] == 0.0


def test_texts_of_the_spheres():
    assert density_text(0, 0) == "empty sphere: r = 0, an enemy at the same point"
    assert density_text(1, 0) == "border object — it can never cover another object"
    assert density_text(3, 1) == "dense; own-class objects on the sphere are not counted"
    assert density_text(2, 0) == "dense: 1 own-class object inside"
    assert relation_meaning(Relation.EQUAL, 3, 3, ZScope.ALL).startswith("e = k")
    assert relation_meaning(Relation.EQUAL, 2, 3, ZScope.ALL).startswith("tie at distance r")
    assert relation_meaning(Relation.OUTSIDE, 1, 3, ZScope.OWN_CLASS).endswith("g ≤ k")
    assert possible_cover(1, 2, LITERAL) == 1  # g = 1: its own sphere cannot put it in L
    assert possible_cover(3, 3, Readings.from_code("211")) == 2  # A = 2
    assert cover_change_text(None, 0, 2, LITERAL) == "moved to D by Step 5 before its turn"
    assert cover_change_text(4, 1, 3, CLOSEST).endswith("or were processed (B = 2)")
    assert cover_change_text(1, 3, 3, LITERAL) == "the same spheres as at the start"
    assert object_title("S₁", "2", 3, 0.0, 1) == "S₁ (class 2)  ·  g = 3  ·  d = 0.000  ·  queue 1"
    assert iteration_title(7, None) == "Iteration 7  ·  queue empty → Step 6 (stop)"


def test_layout_of_the_diagrams(default_run: ExperimentResult):
    c = default_run.result.characteristics
    matrix = default_run.distances.matrix
    layout = radial_layout(matrix, c, 0)
    assert layout.rank.tolist() == [0, 1, 3, 5, 2, 6, 4]
    assert layout.angle[1] == pytest.approx(math.pi / 2)  # the nearest neighbour on top
    assert angle_step(7) == pytest.approx(math.pi / 3)
    x, y = layout.xy(float(matrix.max()))
    assert (x[2], y[2]) == pytest.approx((0.538137796905894, -0.31069400190473))
    assert np.hypot(*layout.xy()) == pytest.approx(matrix[0])
    assert strip_levels(c).tolist() == [7, 4, 6, 1, 5, 2, 3]
    assert ray_length(c) == pytest.approx(8.366366278928)
    assert coinciding_points(c)[0] == (0, 2)  # S₁ and S₃ share (r, z): case 5


# ---------------------------------------------------------------- the core: Steps 2–5


def test_iterations_as_spheres(default_run: ExperimentResult):
    result = default_run.result
    matrix = default_run.distances.matrix
    second = iteration_spheres(result, matrix, 2)
    assert second.selected == 2
    assert second.candidates == (2, 5, 6)
    assert second.covering == (2, 6)
    assert second.members == (2, 6)
    assert np.array_equal(second.sphere_test(), second.in_l)
    assert second.role(6) == f"{Role.IN_L.value}  ·  moves to D (Step 5)"
    assert second.role(5) == Role.F_APART.value
    assert second.role(0) == Role.ENEMY.value
    fourth = iteration_spheres(result, matrix, 4)
    assert fourth.covering_in_d == (0, 4)
    assert fourth.role(0) == Role.D_CONTAINS.value
    fifth = iteration_spheres(result, matrix, 5)
    assert fifth.role(5).startswith(Role.SELECTED_BORDER.value)
    stop = iteration_spheres(result, matrix, 7)
    assert stop.stop
    assert stop.role(2) == Role.FINAL_F.value
    assert stop.role(0) == Role.FINAL_D.value
    with pytest.raises(IndexError):
        iteration_spheres(result, matrix, 0)
    assert cases_in_iteration(default_run.cases, 2) == (9, 10, 11, 13)
    assert set(SHORT_TITLES) == {6, 7, 8, 9, 10, 11, 12, 13, 15}


def test_readings_a_and_b_in_the_sphere_view(default_run: ExperimentResult):
    matrix = default_run.distances.matrix
    y = default_run.dataset.class_index
    a2 = run_erdr(matrix, y, 3, readings=Readings.from_code("211"))
    view = iteration_spheres(a2, matrix, 1)
    assert not view.allowed[view.selected]
    assert view.role(view.selected) == Role.SELECTED_EXCLUDED.value
    b2 = run_erdr(matrix, y, 3, readings=Readings.from_code("121"))
    excluded = [
        (t, j)
        for t in range(1, len(b2.iterations) + 1)
        for j in iteration_spheres(b2, matrix, t).covering
        if not iteration_spheres(b2, matrix, t).in_l[j] and j != b2.iterations[t - 1].selected
    ]
    assert excluded  # B = 2: a processed object's sphere contains Sᵢ but does not count
    t, j = excluded[0]
    assert iteration_spheres(b2, matrix, t).role(j) == Role.EXCLUDED.value


@settings(max_examples=120, deadline=None)
@given(samples(), st.sampled_from(all_readings((ZScope.ALL, ZScope.OWN_CLASS))), st.integers(1, 3))
def test_sphere_invariants_on_random_samples(sample, reading, k):
    D, y = sample
    m = y.size
    sizes = np.bincount(y)
    largest = m - 1 if reading.z_scope is ZScope.ALL else int(sizes.min()) - 1
    if largest < 1:
        return
    c = compute_characteristics(D, y, min(k, largest), z_scope=reading.z_scope)
    p = sphere_profile(D, c)
    assert np.array_equal(c.g, p.g_from_spheres)  # the own class inside, Sᵢ when r > 0
    assert not p.enemies_inside.any()  # no enemy strictly inside: r is the nearest enemy
    assert np.array_equal(p.strictly_inside, c.g)  # closed count − on the sphere
    assert np.abs(p.d_from_ratio - c.d).max() < 1e-9  # d from the ratio of the radii
    if reading.z_scope is ZScope.ALL:  # z < r ⇒ g ≥ k + 1; z > r ⇒ g ≤ k − 1
        for i, relation in enumerate(p.relation):
            if relation is Relation.INSIDE:
                assert c.g[i] >= c.k + 1
            if relation is Relation.OUTSIDE:
                assert c.g[i] <= c.k - 1
    result = run_erdr(D, y, c.k, readings=reading, characteristics=c)
    for it in result.iterations:  # L is exactly the sphere test
        view = iteration_spheres(result, D, it.number)
        assert np.array_equal(view.sphere_test(), view.in_l)
        assert view.members == it.L
        assert set(it.covering) == set(view.covering)


# ---------------------------------------------------------------- tables


def test_sphere_tables(default_run: ExperimentResult):
    tables = {t.key: t for t in run_tables(default_run)}
    for key in ("enemy-spheres", "k-spheres", "covering", "sphere-checks"):
        assert tables[key].sheet == HYPERSPHERES_STEP1
    for key in ("iteration-spheres", "sphere-trace"):
        assert tables[key].sheet == HYPERSPHERES_ITERATIONS
    enemy = enemy_spheres_table(default_run)
    assert enemy.rows[3][-1] == "border object — it can never cover another object"
    assert enemy.mark(0, 9) == "yes"
    k = k_spheres_table(default_run)
    assert k.mark(0, 12) == "selected"  # z = r: the spheres coincide
    covering = covering_table(default_run)
    assert covering.rows[6][6] == "— never selected"
    checks = sphere_checks_table(default_run)
    assert checks.rows[5][2].startswith("✓")
    overview = iteration_spheres_table(default_run)
    assert [row[1] for row in overview.rows] == ["S₁", "S₃", "S₅", "S₂", "S₆", "S₄"]
    assert overview.rows[3][6] == "{S₁, S₅}"  # in D, but their spheres would contain S₂
    trace = sphere_trace_table(default_run)[0]
    assert len(trace.rows) == 6 * 7
    assert trace.matrix
    detail = iteration_detail_table(default_run, 2)
    assert detail.extra["Moved to D"] == "S₇ (Step 5)"
    assert detail.mark(6, 11) == "L"
    assert iteration_detail_table(default_run, 7).rows[0][4] == "—"
    profile = object_spheres_table(default_run, 1)
    assert dict(profile.rows)["L when Sᵢ was selected"] == "{S₂}"
    for table in tables.values():
        if table.sheet.startswith("Hyperspheres"):
            assert stage_of(table.sheet) is not None


def test_large_samples_leave_the_sphere_trace_out(heart270):
    run = run_experiment(dataset=heart270)
    assert sphere_trace_table(run) == []
    assert "sphere-trace" not in {t.key for t in run_tables(run)}
    assert len(iteration_spheres_table(run).rows) == len(run.result.iterations)


# ---------------------------------------------------------------- figures


def test_sphere_figures(default_run: ExperimentResult, tmp_path: Path):
    keys = [key for key, _ in figure_names()]
    assert keys[:5] == ["map", "radii", "priorities", "sizes", "composition"]
    assert {"hyperspheres", "strips", "radius-plane", "iteration-spheres"} <= set(keys)
    assert figure_sheet("map") == "Charts"
    assert figure_sheet("strips") == HYPERSPHERES_STEP1
    assert set(FIGURE_SHEETS) == {"hyperspheres", "strips", "radius-plane", "iteration-spheres"}
    with warnings.catch_warnings():
        warnings.simplefilter("error")  # no layout warning: every figure fits
        for key in PANEL_FIGURES:
            figure = run_figure(key, default_run)
            figure.savefig(tmp_path / f"{key}.png")
        panels = run_figure("hyperspheres", default_run)
        assert sum(ax.has_data() for ax in panels.axes) == 7
        single = object_spheres_figure(default_run, 1)
        single.savefig(tmp_path / "object.png")
        assert single.axes[0].get_title(loc="left").startswith("S₂ (class 2)")
        for t in range(1, 8):
            iteration_spheres_figure(default_run, t).savefig(tmp_path / f"t{t}.png")
    assert sphere_objects(default_run) == list(range(7))
    assert sphere_iterations(default_run) == list(range(1, 7))
    with pytest.raises(KeyError):
        run_figure("nothing", default_run)
    gif = iteration_animation(default_run, tmp_path / "it.gif", dpi=40)
    with Image.open(gif) as image:
        assert image.n_frames == 7  # six iterations and Step 6


def test_sphere_figures_of_a_large_sample(heart270, tmp_path: Path):
    run = run_experiment(dataset=heart270)
    assert run.result.m > CROWDED
    assert len(sphere_objects(run)) == PANEL_LIMIT
    assert sphere_objects(run)[0] == int(run.result.characteristics.priority[0])
    with warnings.catch_warnings():
        warnings.simplefilter("error")
        for figure in (
            run_figure("iteration-spheres", run),
            iteration_spheres_figure(run, 1),
            iteration_spheres_figure(run, len(run.result.iterations) + 1),
            object_spheres_figure(run, 0),
        ):
            figure.savefig(tmp_path / "f.png", dpi=40)


# ---------------------------------------------------------------- exports


def test_exports_place_the_sphere_figures(default_run: ExperimentResult, tmp_path: Path):
    files = export_result(default_run, tmp_path, ["figures", "xlsx", "html", "csv"])
    names = {p.relative_to(tmp_path).as_posix() for p in files}
    assert {"figures/strips.svg", "figures/iteration-spheres.gif"} <= names
    assert "figures/hyperspheres/object-1.png" in names
    assert "figures/iteration-spheres/iteration-6.svg" in names
    assert {"csv/enemy-spheres.csv", "csv/sphere-trace.csv"} <= names
    book = openpyxl.load_workbook(tmp_path / "workbook.xlsx")
    assert book["Hyperspheres (Step 1)"]._images  # the sphere figures beside their tables
    assert len(book["Charts"]._images) == 5  # the other figures, as before
    assert book["Hyperspheres (Iterations)"].sheet_properties.tabColor.rgb.endswith("C55A11")
    html = (tmp_path / "report.html").read_text(encoding="utf-8")
    assert html.index('aria-label="Neighbourhood strips"') < html.index('id="enemy-spheres"')
    assert html.index('aria-label="Hyperspheres in every iteration"') < html.index(
        'id="iteration-spheres"'
    )
    no_gif = export_result(
        default_run, tmp_path / "plain", ["figures"], ExportOptions(figure_formats=("png",))
    )
    assert not any(p.suffix == ".gif" for p in no_gif)


# ---------------------------------------------------------------- the validator


def test_validation_detects_a_changed_sphere_value(tmp_path: Path):
    copy = tmp_path / "workbook.xlsx"
    shutil.copyfile(WORKBOOK, copy)  # the specification itself is never touched
    book = openpyxl.load_workbook(copy, data_only=True)
    book["Hyperspheres (Step 1)"]["E5"] = 99  # own class inside the sphere of S₁
    book["Hypersphere Data"]["F36"] = 0.5  # x of S₂ in the diagram of S₁
    book["Hyperspheres (Iterations)"]["O24"] = "something else"  # a role in iteration 1
    book.save(copy)
    report = validate_workbook(copy, skip=("table1",))
    failed = {(check.section, check.name) for check in report.failures}
    assert failed == {
        ("Hyperspheres (Step 1)", "enemy sphere of S₁ · own class inside"),
        ("Hypersphere Data", "diagram of S₁ · S₂ · x"),
        ("Hyperspheres (Iterations)", "iteration 1 · S₁ · role"),
    }
