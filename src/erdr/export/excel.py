"""The run as an Excel workbook laid out like the ERDR workbook (ADR-012).

Tables that mirror the same sheet of the ERDR workbook are stacked on one sheet, in workbook
order (*Overview*, *Dataset*, *Normalized Dataset*, *Zhuravlev Distances*, …); a figure that
belongs to a sheet (the hypersphere figures) goes under that sheet's tables, the others to
*Charts*. Values are written exactly (no rounding). The styles are the workbook's: dark-blue
sheet title, blue table title bars, grey-blue column headers, sheet tabs coloured by stage; the
marks of the cells and headers become the colours of :mod:`erdr.export.theme`, explained by the
colour legend on the *Overview* sheet (ADR-019).
"""

from __future__ import annotations

from collections.abc import Sequence
from io import BytesIO
from pathlib import Path

from matplotlib.figure import Figure
from openpyxl import Workbook
from openpyxl.cell.cell import Cell as XlCell
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from erdr import __version__
from erdr.export.tables import Table
from erdr.export.theme import legend, meaning, stage_of, swatch

SHEET_TITLE = "1F3864"
"""The workbook's dark blue: the title bar of the *Overview* sheet."""
TABLE_TITLE = "2F5597"
"""The workbook's blue: the title bar of every table."""
HEADER = "D6DCE4"
"""The workbook's grey-blue column header (a column without a mark)."""
NOTE_INK = "595959"
_WIDTH_LIMIT = 60


def _rgb(colour: str) -> str:
    """``#ddebf7`` → ``DDEBF7`` (openpyxl's form)."""
    return colour.lstrip("#").upper()


def _paint(cell: XlCell, mark: str | None, *, bold: bool = False) -> None:
    """Give ``cell`` the fill and ink of ``mark`` (nothing for no mark)."""
    found = swatch(mark)
    if found is None:
        if bold:
            cell.font = Font(bold=True)
        return
    cell.fill = PatternFill("solid", fgColor=_rgb(found.fill))
    cell.font = Font(bold=bold, color=_rgb(found.ink))


def _sheet_name(name: str, taken: set[str]) -> str:
    base = "".join("_" if ch in "[]:*?/\\" else ch for ch in name)[:31] or "Sheet"
    candidate, counter = base, 2
    while candidate in taken:
        suffix = f" ({counter})"
        candidate = base[: 31 - len(suffix)] + suffix
        counter += 1
    taken.add(candidate)
    return candidate


def _write_table(ws: Worksheet, table: Table, top: int) -> int:
    """Write ``table`` from row ``top``; returns the first free row after it."""
    width = len(table.columns)
    title = ws.cell(row=top, column=1, value=table.title)
    title.font = Font(bold=True, color="FFFFFF")
    for column in range(1, width + 1):
        ws.cell(row=top, column=column).fill = PatternFill("solid", fgColor=TABLE_TITLE)
    for column, name in enumerate(table.columns, start=1):
        cell = ws.cell(row=top + 1, column=column, value=name)
        mark = table.header_mark(column - 1)
        if swatch(mark) is None:
            cell.fill = PatternFill("solid", fgColor=HEADER)
        _paint(cell, mark, bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r, row in enumerate(table.rows):
        for c, value in enumerate(row):
            cell = ws.cell(row=top + 2 + r, column=c + 1, value=value)
            _paint(cell, table.mark(r, c), bold=c == 0)
    end = top + 2 + len(table.rows)
    if table.note:
        ws.cell(row=end, column=1, value=table.note).font = Font(italic=True, color=NOTE_INK)
        end += 1
    return end + 1


FigureItem = tuple[str, Figure] | tuple[str, Figure, str]
"""A figure for the workbook: (name, figure) or (name, figure, sheet it belongs to)."""


def _insert_figure(ws: Worksheet, name: str, figure: Figure, top: int) -> int:
    """Put ``figure`` with its name from row ``top``; returns the first free row after it."""
    ws.cell(row=top, column=1, value=name).font = Font(bold=True)
    buffer = BytesIO()
    figure.savefig(buffer, format="png", dpi=110)
    buffer.seek(0)
    image = Image(buffer)
    ws.add_image(image, f"A{top + 1}")
    return top + int(image.height / 20) + 3


def _fit_columns(ws: Worksheet) -> None:
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None or cell.row == 1:
                continue
            text = str(cell.value)
            widths[cell.column] = max(widths.get(cell.column, 0), min(len(text), _WIDTH_LIMIT))
    for column, width in widths.items():
        ws.column_dimensions[get_column_letter(column)].width = max(8, width + 2)


def _tab(ws: Worksheet, sheet: str) -> None:
    """Colour the sheet tab by the stage of the experiment (the workbook's tab colours)."""
    stage = stage_of(sheet)
    if stage is not None:
        ws.sheet_properties.tabColor = _rgb(stage.colour)


def _legend(
    ws: Worksheet, row: int, tables: Sequence[Table], class_labels: Sequence[object]
) -> int:
    """The colour legend of the marks the workbook uses; returns the next free row."""
    title = ws.cell(row=row, column=1, value="Colour legend")
    title.font = Font(bold=True, color=SHEET_TITLE)
    used: set[str] = set().union(*(table.used_marks() for table in tables))
    for section, marks in legend(used):
        row += 1
        ws.cell(row=row, column=1, value=section).font = Font(italic=True, color=NOTE_INK)
        for mark in marks:
            row += 1
            _paint(ws.cell(row=row, column=1, value=meaning(mark, class_labels)), mark, bold=True)
    return row + 1


def write_workbook(
    path: Path,
    tables: Sequence[Table],
    *,
    title: str,
    settings: Sequence[tuple[str, str]] = (),
    figures: Sequence[FigureItem] = (),
    class_labels: Sequence[object] = (),
) -> Path:
    """Write the tables (stacked per workbook sheet), the settings and the figures.

    Args:
        path: The ``.xlsx`` file.
        tables: The tables; those with the same :attr:`Table.sheet` share a sheet.
        title: The title on the *Overview* sheet.
        settings: (name, value) pairs shown on the *Overview* sheet; a pair named ``Note`` is a
            reading that is not the literal one and is shown as a warning.
        figures: (name, figure) pairs placed on the *Charts* sheet, or (name, figure, sheet)
            placed under the tables of ``sheet`` (on *Charts* when no table is on it).
        class_labels: The labels of the classes K1, K2, … (named in the colour legend).
    """
    workbook = Workbook()
    overview = workbook.active
    assert overview is not None  # a new workbook always has its first sheet
    overview.title = "Overview"
    _tab(overview, "Overview")
    heading = overview.cell(row=1, column=1, value=title)
    heading.font = Font(bold=True, size=14, color="FFFFFF")
    for column in (1, 2):
        overview.cell(row=1, column=column).fill = PatternFill("solid", fgColor=SHEET_TITLE)
    overview.cell(row=2, column=1, value=f"Written by erdr {__version__}").font = Font(
        italic=True, color=NOTE_INK
    )
    row = 4
    for name, value in settings:
        mark = "warning" if name == "Note" else "parameter"
        _paint(overview.cell(row=row, column=1, value=name), mark, bold=True)
        value_cell = overview.cell(row=row, column=2, value=value)
        if mark == "warning":
            _paint(value_cell, mark)
        row += 1
    row += 1
    taken = {"Overview"}
    sheets: dict[str, list[Table]] = {}
    for table in tables:
        sheets.setdefault(table.sheet or table.key, []).append(table)
    for table in sheets.pop("Overview", []):
        row = _write_table(overview, table, row)
    _legend(overview, row, tables, class_labels)
    _fit_columns(overview)

    placed: dict[str, list[tuple[str, Figure]]] = {}
    charted: list[tuple[str, Figure]] = []
    for item in figures:
        sheet = item[2] if len(item) == 3 else None
        if sheet is not None and sheet in sheets:
            placed.setdefault(sheet, []).append((item[0], item[1]))
        else:
            charted.append((item[0], item[1]))
    for name, group in sheets.items():
        ws = workbook.create_sheet(_sheet_name(name, taken))
        _tab(ws, name)
        top = 1
        for table in group:
            top = _write_table(ws, table, top)
        if len(group) == 1:
            ws.freeze_panes = ws.cell(row=3, column=2)
        _fit_columns(ws)
        for figure_name, figure in placed.get(name, []):
            top = _insert_figure(ws, figure_name, figure, top)

    if charted:
        charts = workbook.create_sheet(_sheet_name("Charts", taken))
        _tab(charts, "Charts")
        top = 1
        for name, figure in charted:
            top = _insert_figure(charts, name, figure, top)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path
