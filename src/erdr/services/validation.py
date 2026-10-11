"""Validation of the package against the ERDR workbook (``erdr validate``, ADR-001).

The workbook ``ERDR algorithm [Heart-Disease (7, 13, 2)].xlsx`` is the executable specification.
Its inputs — the 7 selected objects, k, ε and the readings A, B, C — are read from the workbook
itself, the engine runs, and every computed sheet is compared cell by cell with the cached cell
values: the normalized data, the three distance matrices, the neighbour ranks, r, g, z, d and the
enemy spheres, the queue with its tie criteria, every iteration of the trace and the counters,
the routes of the objects, the 17 cases, all eight readings, the 270-object comparison with the
article's Table 1, and finally the workbook's own *Validation* sheet.
"""

from __future__ import annotations

import logging
import math
import warnings
from collections.abc import Callable, Iterable, Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import numpy as np

from erdr.config.models import ERDRConfig, ExperimentConfig, ReadingsConfig
from erdr.core.algorithm import ERDRResult, run_erdr
from erdr.core.readings import CLOSEST, LITERAL, Readings, ZScope, all_readings
from erdr.data.loaders import load_dataset
from erdr.data.schema import Dataset, FeatureType
from erdr.errors import DatasetError
from erdr.notation import DASH, changes_text, moves_text, object_order, object_set
from erdr.services.experiment import ExperimentResult, run_experiment, step0

logger = logging.getLogger(__name__)

DEFAULT_TOLERANCE = 1e-9
"""Absolute tolerance of numeric checks — the workbook's own (Validation sheet)."""
WORKBOOK_NAME = "ERDR algorithm [Heart-Disease (7, 13, 2)].xlsx"
"""File name of the specification workbook."""

_DECIDED = {"d": "d", "r": "r (equal d)", "index": "index (equal d and r)", "—": "— (last)"}
_ENTERED_BY = {4: "Step 4 · |L| = 1", 5: "Step 5 · g fell to 1", None: "— (stays in F)"}
_TABLE1_READINGS: tuple[tuple[str, Readings, int], ...] = (
    ("literal reading", LITERAL, 3),
    ("B = 2", Readings.from_code("121"), 6),
    ("own-class z", Readings.from_code("111", ZScope.OWN_CLASS), 9),
    ("B = 2 + own-class z", CLOSEST, 12),
    ("A = 2", Readings.from_code("211"), 15),
)
"""The five readings of *Article Table 1 Check*: name, reading, 0-based column of |D|."""


@dataclass(frozen=True)
class Check:
    """One compared value."""

    section: str
    """The workbook sheet."""
    name: str
    """What is compared, e.g. ``r(S₁)``."""
    expected: object
    """The workbook's cached value."""
    actual: object
    """The package's value."""
    passed: bool
    """Within the tolerance (numbers) or equal (text, integers)."""


@dataclass(frozen=True)
class ValidationReport:
    """Every check of one workbook."""

    workbook: Path
    tolerance: float
    checks: tuple[Check, ...]

    @property
    def passed(self) -> bool:
        """Whether every check passed."""
        return all(check.passed for check in self.checks)

    @property
    def failures(self) -> tuple[Check, ...]:
        """The checks that failed."""
        return tuple(check for check in self.checks if not check.passed)

    def sections(self) -> list[tuple[str, int, int]]:
        """(sheet, checks, passed) per sheet, in workbook order."""
        order: list[str] = []
        counts: dict[str, list[int]] = {}
        for check in self.checks:
            if check.section not in counts:
                order.append(check.section)
                counts[check.section] = [0, 0]
            counts[check.section][0] += 1
            counts[check.section][1] += int(check.passed)
        return [(name, counts[name][0], counts[name][1]) for name in order]


class _Sheets:
    """Cached cell values of the workbook's sheets."""

    def __init__(self, path: Path) -> None:
        import openpyxl

        try:
            with warnings.catch_warnings():
                warnings.simplefilter("ignore", UserWarning)
                self._book = openpyxl.load_workbook(path, data_only=True)
        except Exception as error:  # openpyxl raises many types for damaged or foreign files
            raise DatasetError(f"{path.name}: cannot read the workbook ({error})") from error
        required = (
            "Parameters",
            "Dataset",
            "Normalized Dataset",
            "Zhuravlev Distances",
            "ERDR Characteristics",
            "ERDR Trace",
        )
        missing = [name for name in required if name not in self._book.sheetnames]
        if missing:
            raise DatasetError(f"{path.name}: not an ERDR workbook (no sheet {', '.join(missing)})")

    def has(self, sheet: str) -> bool:
        return sheet in self._book.sheetnames

    def cell(self, sheet: str, row: int, column: int) -> Any:
        """Value of the cell at 1-based (row, column)."""
        return self._book[sheet].cell(row=row, column=column).value

    def close(self) -> None:
        self._book.close()


class _Collector:
    def __init__(self, tolerance: float) -> None:
        self.tolerance = tolerance
        self.checks: list[Check] = []

    def number(self, section: str, name: str, expected: Any, actual: float) -> None:
        try:
            value = float(expected)
            ok = math.isclose(value, float(actual), rel_tol=0.0, abs_tol=self.tolerance)
        except (TypeError, ValueError):
            ok = False
        self.checks.append(Check(section, name, expected, actual, ok))

    def exact(self, section: str, name: str, expected: Any, actual: object) -> None:
        if isinstance(expected, float) and expected.is_integer():
            expected = int(expected)
        ok = expected == actual or (
            isinstance(expected, str) and isinstance(actual, str) and expected.strip() == actual
        )
        self.checks.append(Check(section, name, expected, actual, ok))


def _workbook_config(sheets: _Sheets) -> ExperimentConfig:
    """k, ε and the readings A, B, C from the sheet *Parameters*."""
    code = "".join(str(int(sheets.cell("Parameters", row, 2))) for row in (24, 25, 26))
    return ExperimentConfig(
        erdr=ERDRConfig(
            k=int(sheets.cell("Parameters", 12, 2)),
            epsilon=float(sheets.cell("Parameters", 17, 2)),
            k_rule="warn",
            readings=ReadingsConfig.from_readings(Readings.from_code(code)),
        )
    )


def _source_dataset(sheets: _Sheets) -> Dataset | None:
    """Heart-Disease (270, 13, 2) from the sheet *Source Dataset (270)* (rows 5–274)."""
    sheet = "Source Dataset (270)"
    if not sheets.has(sheet):
        return None
    rows = range(5, 275)
    X = np.array([[float(sheets.cell(sheet, r, c)) for c in range(2, 15)] for r in rows])
    labels = np.array([int(sheets.cell(sheet, r, 15)) for r in rows], dtype=np.int64)
    types = tuple(FeatureType.from_flag(int(sheets.cell(sheet, 275, c))) for c in range(2, 15))
    return Dataset(X=X, y=labels, feature_types=types, name="Heart-Disease (270, 13, 2)")


def _names(dataset: Dataset) -> Sequence[str]:
    return dataset.object_ids


def _check_inputs(sheets: _Sheets, run: ExperimentResult, out: _Collector) -> None:
    section = "Parameters"
    sizes = run.dataset.class_sizes
    out.exact(section, "N", sheets.cell(section, 5, 2), run.dataset.m)
    out.exact(section, "|K1|", sheets.cell(section, 6, 2), sizes[0])
    out.exact(section, "|K2|", sheets.cell(section, 7, 2), sizes[1])
    out.exact(section, "k_max", sheets.cell(section, 10, 2), run.permitted.k_max)
    out.exact(section, "permitted k", str(sheets.cell(section, 11, 2)), run.permitted.text)

    source = _source_dataset(sheets)
    if source is not None:
        rows = [int(sheets.cell("Dataset", 5 + i, 2)) for i in range(run.dataset.m)]
        for i, row in enumerate(rows):
            same = bool(np.array_equal(source.X[row - 1], run.dataset.X[i]))
            out.exact("Dataset", f"{_names(run.dataset)[i]} = source row {row}", True, same)


def _check_step0(sheets: _Sheets, run: ExperimentResult, out: _Collector) -> None:
    names = _names(run.dataset)
    m, n = run.dataset.m, run.dataset.n
    section = "Normalized Dataset"
    for i in range(m):
        for j in range(n):
            out.number(
                section,
                f"x′({names[i]}, x{j + 1})",
                sheets.cell(section, 5 + i, 2 + j),
                run.distances.unified[i, j],
            )
    stats = run.distances.scaling.statistics
    for feature in (int(f) for f in np.flatnonzero(run.dataset.quantitative)):
        column = 2 + feature
        name = f"x{feature + 1}"
        out.number(section, f"min {name}", sheets.cell(section, 13, column), stats["min"][feature])
        out.number(section, f"max {name}", sheets.cell(section, 14, column), stats["max"][feature])

    section = "Zhuravlev Distances"
    blocks: list[tuple[str, int, np.ndarray | None]] = [
        ("ρ", 5, run.distances.matrix),
        ("ρ_I", 15, run.distances.quantitative_part),
        ("ρ_J", 25, run.distances.nominal_part),
    ]
    for label, top, matrix in blocks:
        if matrix is None:
            continue
        for i in range(m):
            for j in range(m):
                out.number(
                    section,
                    f"{label}({names[i]}, {names[j]})",
                    sheets.cell(section, top + i, 2 + j),
                    matrix[i, j],
                )


def _check_step1(sheets: _Sheets, run: ExperimentResult, out: _Collector) -> None:
    names = _names(run.dataset)
    c = run.result.characteristics
    m = c.m
    if sheets.has("Sorted Neighbors"):
        ranks = np.zeros((m, m), dtype=np.int64)
        for i in range(m):
            ranks[i, c.order[i]] = np.arange(1, m)
        for i in range(m):
            for j in range(m):
                if i != j:
                    out.exact(
                        "Sorted Neighbors",
                        f"rank of {names[j]} for {names[i]}",
                        sheets.cell("Sorted Neighbors", 5 + i, 2 + j),
                        int(ranks[i, j]),
                    )
    section = "ERDR Characteristics"
    for i in range(m):
        row = 5 + i
        out.exact(
            section,
            f"nearest enemy of {names[i]}",
            sheets.cell(section, row, 3),
            names[c.nearest_enemy[i]],
        )
        out.number(section, f"r({names[i]})", sheets.cell(section, row, 4), c.r[i])
        out.exact(section, f"g({names[i]})", sheets.cell(section, row, 5), int(c.g[i]))
        out.exact(
            section,
            f"rank of the nearest enemy of {names[i]}",
            sheets.cell(section, row, 6),
            int(c.enemy_rank[i]),
        )
        out.exact(
            section,
            f"k-th neighbour of {names[i]}",
            sheets.cell(section, row, 7),
            names[c.kth_neighbour[i]],
        )
        out.number(section, f"z({names[i]}, {c.k})", sheets.cell(section, row, 8), c.z[i])
        out.number(section, f"z − r of {names[i]}", sheets.cell(section, row, 9), c.z[i] - c.r[i])
        out.number(section, f"d({names[i]})", sheets.cell(section, row, 10), c.d[i])
        out.exact(
            section,
            f"priority position of {names[i]}",
            sheets.cell(section, row, 11),
            int(c.position[i]),
        )
        inside = [j for j in c.order[i] if c.spheres[i, j]]
        listed = ", ".join([f"{names[i]} (itself)", *(names[j] for j in inside)])
        out.exact(
            section,
            f"own class inside the sphere of {names[i]}",
            sheets.cell(section, row, 12),
            listed,
        )
        for j in range(m):
            out.exact(
                section,
                f"M({names[i]}, {names[j]})",
                sheets.cell(section, 16 + i, 2 + j),
                int(c.spheres[i, j]),
            )
    if sheets.has("Priority Queue"):
        decided = c.decided_by()
        for position, i in enumerate(c.priority):
            row = 5 + position
            out.exact(
                "Priority Queue",
                f"position {position + 1}",
                sheets.cell("Priority Queue", row, 2),
                names[i],
            )
            out.exact(
                "Priority Queue",
                f"position {position + 1} decided by",
                sheets.cell("Priority Queue", row, 7),
                _DECIDED[decided[position]],
            )


def _check_trace(sheets: _Sheets, run: ExperimentResult, out: _Collector) -> None:
    names = _names(run.dataset)
    result = run.result
    m = result.m
    section = "ERDR Trace"
    for it in result.iterations:
        row = 4 + it.number
        t = f"iteration {it.number}"
        out.exact(section, f"{t} · selected", sheets.cell(section, row, 2), names[it.selected])
        out.number(section, f"{t} · d(Sᵢ)", sheets.cell(section, row, 3), it.d)
        out.number(section, f"{t} · r(Sᵢ)", sheets.cell(section, row, 4), it.r)
        out.exact(section, f"{t} · |L|", sheets.cell(section, row, 5), len(it.L))
        out.exact(section, f"{t} · |F| after", sheets.cell(section, row, 6), it.f_size)
        out.exact(section, f"{t} · |D| after", sheets.cell(section, row, 7), it.d_size)
        out.exact(section, f"{t} · queue left", sheets.cell(section, row, 8), it.queue_left)
        out.exact(section, f"{t} · L", sheets.cell(section, row, 9), object_set(it.L, names))
        out.exact(section, f"{t} · decision", sheets.cell(section, row, 10), it.decision.value)
        changes = [(ch.object, ch.before, ch.after) for ch in it.changes]
        out.exact(
            section, f"{t} · counters", sheets.cell(section, row, 11), changes_text(changes, names)
        )
        moves = [(mv.object, mv.step) for mv in it.moves]
        out.exact(
            section, f"{t} · moved to D", sheets.cell(section, row, 12), moves_text(moves, names)
        )
    stop_row = 5 + len(result.iterations)
    out.exact(section, "stop after the last iteration", sheets.cell(section, stop_row, 2), "— stop")
    history = result.counter_history()
    blocks = m  # the workbook has one block per object (at most m iterations)
    for block in range(blocks + 1):
        counters = history[min(block, len(result.iterations))]
        in_f = result.state_after(min(block, len(result.iterations))).in_f
        label = "start" if block == 0 else f"after iteration {block}"
        for j in range(m):
            out.exact(
                section,
                f"g({names[j]}) {label}",
                sheets.cell(section, 77 + block, 2 + j),
                int(counters[j]),
            )
            out.exact(
                section,
                f"{names[j]} in F {label}",
                sheets.cell(section, 66 + block, 2 + j),
                int(in_f[j]),
            )
    for j, route in enumerate(result.routes):
        out.exact(
            section,
            f"{names[j]} processed in",
            sheets.cell(section, 88, 2 + j),
            route.processed_in or DASH,
        )
        out.exact(
            section,
            f"{names[j]} moved to D in",
            sheets.cell(section, 89, 2 + j),
            route.entered_d_in or DASH,
        )
        by = f"Step {route.entered_by}" if route.entered_by else DASH
        out.exact(section, f"{names[j]} moved by", sheets.cell(section, 90, 2 + j), by)
        out.exact(
            section,
            f"{names[j]} counter decreases",
            sheets.cell(section, 91, 2 + j),
            route.counter_decreases,
        )
        out.exact(
            section,
            f"{names[j]} final sample",
            sheets.cell(section, 92, 2 + j),
            "D" if route.in_training else "F",
        )


def _check_result(sheets: _Sheets, run: ExperimentResult, out: _Collector) -> None:
    names = _names(run.dataset)
    result = run.result
    section = "Training & Validation"
    if not sheets.has(section):
        return
    for j, route in enumerate(result.routes):
        row = 5 + j
        sample = "D · training" if route.in_training else "F · validation"
        out.exact(section, f"{names[j]} sample", sheets.cell(section, row, 4), sample)
        out.exact(
            section,
            f"{names[j]} processed in",
            sheets.cell(section, row, 5),
            route.processed_in or DASH,
        )
        out.exact(
            section,
            f"{names[j]} entered D in",
            sheets.cell(section, row, 6),
            route.entered_d_in or DASH,
        )
        out.exact(
            section,
            f"{names[j]} entered D by",
            sheets.cell(section, row, 7),
            _ENTERED_BY[route.entered_by],
        )
        out.exact(
            section,
            f"{names[j]} counter decreases",
            sheets.cell(section, row, 8),
            route.counter_decreases,
        )
    for row, label, indices in ((15, "D", result.training), (16, "F", result.validation)):
        counts = run.class_counts(indices)
        out.exact(section, f"|{label} ∩ K1|", sheets.cell(section, row, 2), counts[0])
        out.exact(section, f"|{label} ∩ K2|", sheets.cell(section, row, 3), counts[1])
        out.exact(section, f"|{label}|", sheets.cell(section, row, 4), int(indices.size))
        out.number(section, f"|{label}| / N", sheets.cell(section, row, 5), indices.size / result.m)
    out.exact(section, "D", sheets.cell(section, 19, 2), object_set(result.training, names))
    out.exact(section, "F", sheets.cell(section, 20, 2), object_set(result.validation, names))

    section = "Case Checklist"
    if sheets.has(section):
        for case in run.cases:
            row = 4 + case.case.number
            out.exact(
                section,
                f"case {case.case.number} occurs",
                sheets.cell(section, row, 4),
                int(case.occurs),
            )
            first = sheets.cell(section, row, 5)
            actual: object = int(case.first) if case.first.isdigit() else case.first
            out.exact(section, f"case {case.case.number} first", first, actual)


def _run_reading(base: ExperimentResult, reading: Readings) -> ERDRResult:
    return run_erdr(
        base.distances.matrix,
        base.dataset.class_index,
        base.k,
        epsilon=base.result.epsilon,
        readings=reading,
    )


def _check_readings(sheets: _Sheets, run: ExperimentResult, out: _Collector) -> None:
    section = "Readings (Switches)"
    if not sheets.has(section):
        return
    names = _names(run.dataset)
    for offset, reading in enumerate(all_readings()):
        row = 11 + offset
        code = "".join(str(int(sheets.cell(section, row, col))) for col in (1, 2, 3))
        out.exact(section, f"row {row} reading", code, reading.code)
        result = _run_reading(run, reading)
        label = f"reading {reading.code}"
        out.exact(
            section,
            f"{label} · order",
            sheets.cell(section, row, 4),
            object_order(result.processing_order, names),
        )
        out.exact(
            section,
            f"{label} · D",
            sheets.cell(section, row, 5),
            object_set(result.training, names),
        )
        out.exact(
            section,
            f"{label} · F",
            sheets.cell(section, row, 6),
            object_set(result.validation, names),
        )
        counts = run.class_counts(result.training)
        out.exact(
            section, f"{label} · |D|", sheets.cell(section, row, 7), int(result.training.size)
        )
        out.exact(section, f"{label} · |D ∩ K1|", sheets.cell(section, row, 8), counts[0])
        out.exact(section, f"{label} · |D ∩ K2|", sheets.cell(section, row, 9), counts[1])
        out.exact(
            section, f"{label} · iterations", sheets.cell(section, row, 10), len(result.iterations)
        )


def _check_table1(sheets: _Sheets, run: ExperimentResult, out: _Collector) -> None:
    section = "Article Table 1 Check"
    source = _source_dataset(sheets)
    if source is None or not sheets.has(section):
        return
    matrix = step0(run.config, source).matrix
    classes = source.class_index

    def sizes(k: int, reading: Readings) -> tuple[int, int, int]:
        result = run_erdr(matrix, classes, k, epsilon=run.result.epsilon, readings=reading)
        counts = np.bincount(classes[result.training], minlength=2)
        return int(result.training.size), int(counts[0]), int(counts[1])

    for row in range(6, 11):
        k = int(sheets.cell(section, row, 1))
        for label, reading, column in _TABLE1_READINGS:
            size, k1, k2 = sizes(k, reading)
            out.exact(
                section, f"k = {k} · {label} · |D|", sheets.cell(section, row, column + 1), size
            )
            out.exact(
                section,
                f"k = {k} · {label} · K1, K2",
                sheets.cell(section, row, column + 2),
                f"{k1}, {k2}",
            )
    for row in range(30, 50):
        k = int(sheets.cell(section, row, 1))
        out.exact(
            section, f"k = {k} · |D| literal", sheets.cell(section, row, 2), sizes(k, LITERAL)[0]
        )
        out.exact(
            section, f"k = {k} · |D| closest", sheets.cell(section, row, 4), sizes(k, CLOSEST)[0]
        )


def _validation_sheet_values(run: ExperimentResult) -> dict[str, object]:
    """The engine's value for every check name of the workbook's *Validation* sheet."""
    names = _names(run.dataset)
    result = run.result
    c = result.characteristics
    values: dict[str, object] = {}
    for i, name in enumerate(names):
        values[f"r({name})"] = c.r[i]
        values[f"g({name})"] = int(c.g[i])
        values[f"z({name}, {c.k})"] = c.z[i]
        values[f"d({name})"] = c.d[i]
    values["Step 2 order"] = object_order(c.priority, names)
    for t in range(1, result.m + 1):
        it = result.iterations[t - 1] if t <= len(result.iterations) else None
        values[f"Iteration {t} · selected"] = (
            names[it.selected] if it else "— (queue empty → Step 6)"
        )
        values[f"Iteration {t} · |L|"] = len(it.L) if it else 0
        values[f"Iteration {t} · moved to D"] = (
            moves_text([(mv.object, mv.step) for mv in it.moves], names) if it else DASH
        )
    values["D · training sample"] = object_set(result.training, names)
    values["F · validation sample"] = object_set(result.validation, names)
    values["Cases that occur (of 17)"] = sum(case.occurs for case in run.cases)
    return values


def _check_validation_sheet(sheets: _Sheets, run: ExperimentResult, out: _Collector) -> None:
    section = "Validation"
    if not sheets.has(section):
        return
    engine = _validation_sheet_values(run)
    row = 6
    while (name := sheets.cell(section, row, 1)) not in (None, "All checks"):
        expected = sheets.cell(section, row, 2)
        key = str(name).strip()
        if key in engine:
            actual = engine[key]
            if isinstance(actual, float):
                out.number(section, key, expected, actual)
            else:
                out.exact(section, key, expected, actual)
        else:
            out.exact(section, key, expected, "(no engine value)")
        row += 1


_STEPS: tuple[Callable[[_Sheets, ExperimentResult, _Collector], None], ...] = (
    _check_inputs,
    _check_step0,
    _check_step1,
    _check_trace,
    _check_result,
    _check_readings,
    _check_table1,
    _check_validation_sheet,
)


def validate_workbook(
    path: Path,
    *,
    tolerance: float = DEFAULT_TOLERANCE,
    skip: Iterable[str] = (),
) -> ValidationReport:
    """Compare the package with every computed sheet of the ERDR workbook.

    Args:
        path: The workbook (``ERDR algorithm [Heart-Disease (7, 13, 2)].xlsx``).
        tolerance: Absolute tolerance of numeric checks.
        skip: Names of check groups to leave out (``table1`` skips the 270-object runs).

    Raises:
        DatasetError: The file is not a readable ERDR workbook.
    """
    sheets = _Sheets(path)
    try:
        dataset = load_dataset(path, "erdr-workbook")
        run = run_experiment(_workbook_config(sheets), dataset)
        out = _Collector(tolerance)
        skipped = set(skip)
        for step in _STEPS:
            if step.__name__.removeprefix("_check_") in skipped:
                continue
            step(sheets, run, out)
    finally:
        sheets.close()
    report = ValidationReport(path, tolerance, tuple(out.checks))
    logger.info(
        "workbook validated",
        extra={"checks": len(report.checks), "failed": len(report.failures)},
    )
    return report


def find_workbook(start: Path | None = None) -> Path | None:
    """``resources/Experiments/<WORKBOOK_NAME>`` in the project above ``start``, if present."""
    folder = (start or Path.cwd()).resolve()
    for candidate in (folder, *folder.parents):
        path = candidate / "resources" / "Experiments" / WORKBOOK_NAME
        if path.is_file():
            return path
    return None
